'''
Multi-format support (PDF, DOCX, TXT, MD, HTML, tabular, code, notebooks, images)
- Intelligent chunking with overlap
- Metadata extraction (title, author, date, file type)
- Text cleaning and normalization
- Duplicate detection
'''
from __future__ import annotations

import hashlib
import logging
import os
import re
from datetime import datetime
from typing import Any, Dict, List, Optional, Protocol

import frontmatter
import tiktoken
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from src.core.extractors import extract_for_path
from src.core.supported_formats import is_supported

logger = logging.getLogger(__name__)

_CODE_FENCE_RE = re.compile(r"```[\s\S]*?```", re.MULTILINE)
_WIKILINK_RE = re.compile(r"\[\[([^\]]+)\]\]")


class _RegexTokenizer:
    """Offline fallback tokenizer when tiktoken encoding cannot be loaded."""

    _token_pattern = re.compile(r"\w+|[^\w\s]", re.UNICODE)

    def encode(self, text: str) -> List[str]:
        return self._token_pattern.findall(text)

    def decode(self, token_ids: List[str]) -> str:
        if not token_ids:
            return ""
        return " ".join(token_ids)


class _Tokenizer(Protocol):
    def encode(self, text: str) -> List[Any]:
        ...

    def decode(self, token_ids: List[Any]) -> str:
        ...


class DocumentProcessor:
    def __init__(self, chunk_size: int = 600, overlap: int = 100, tokenizer_name: str = "gpt2"):
        if chunk_size <= 0:
            raise ValueError("chunk_size must be > 0")
        if overlap < 0:
            raise ValueError("overlap must be >= 0")
        if overlap >= chunk_size:
            raise ValueError("overlap must be smaller than chunk_size")
        self.chunk_size = chunk_size
        self.overlap = overlap
        self.tokenizer_name = tokenizer_name
        self._tokenizer: _Tokenizer
        try:
            self._tokenizer = tiktoken.get_encoding(tokenizer_name)
        except Exception:
            self._tokenizer = _RegexTokenizer()
        self._seen_hashes: set = set()

    @staticmethod
    def strip_code_fences(text: str) -> str:
        """Replace fenced code blocks with whitespace so chunking skips them."""
        return _CODE_FENCE_RE.sub(" ", text)

    @staticmethod
    def resolve_wikilinks(text: str, vault_root: str) -> str:
        """Best-effort: replace [[note]] with wiki-relative path if file exists."""

        def repl(match: re.Match[str]) -> str:
            inner = match.group(1)
            target = inner.split("|", 1)[0].strip()
            wiki = os.path.join(vault_root, "wiki")
            for name in (target, f"{target}.md", f"{target.replace(' ', '-')}.md"):
                cand = os.path.join(wiki, name)
                if os.path.isfile(cand):
                    return os.path.relpath(cand, vault_root).replace("\\", "/")
            return match.group(0)

        return _WIKILINK_RE.sub(repl, text)

    def process_document(
        self,
        file_path: str,
        *,
        vault_root: Optional[str] = None,
        relpath: Optional[str] = None,
        vision_enabled: bool = False,
        vision_max_bytes: int = 5_000_000,
        vision_max_edge: int = 4096,
    ) -> Optional[Dict]:
        ext = os.path.splitext(file_path)[1].lower()
        metadata = self.extract_metadata(file_path)
        if relpath:
            metadata["relpath"] = relpath.replace("\\", "/").lstrip("/")
        if ext == ".md":
            with open(file_path, encoding="utf-8") as f:
                post = frontmatter.load(f)
            fm = dict(post.metadata or {})
            raw_body = post.content or ""
            if fm.get("title"):
                metadata["title"] = str(fm["title"])
            tags = fm.get("tags")
            if isinstance(tags, str):
                tags = [tags]
            elif tags is None:
                tags = []
            metadata["tags"] = list(tags)
            metadata["section"] = str(fm.get("section") or metadata.get("section") or "")
            doc_type = fm.get("type")
            metadata["doc_type"] = str(doc_type) if doc_type is not None else ""
            text = raw_body
        else:
            text = self.extract_text(
                file_path,
                vision_enabled=vision_enabled,
                vision_max_bytes=vision_max_bytes,
                vision_max_edge=vision_max_edge,
            )

        if vault_root:
            text = self.resolve_wikilinks(text, vault_root)

        if self._is_duplicate(text, file_path):
            return None

        text = self.strip_code_fences(text)
        cleaned_text = self.clean_text(text)
        chunks = self.chunk_text(cleaned_text)
        return {
            'metadata': metadata,
            'chunks': chunks,
        }

    def _is_duplicate(self, text: str, file_path: str) -> bool:
        """Skip only when the same absolute path was already processed with this body.

        Different paths with identical content are allowed so shared templates do not
        silently drop later files during a vault scan.
        """
        digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
        key = (os.path.abspath(file_path), digest)
        if key in self._seen_hashes:
            return True
        self._seen_hashes.add(key)
        return False

    def extract_text(
        self,
        file_path: str,
        *,
        vision_enabled: bool = False,
        vision_max_bytes: int = 5_000_000,
        vision_max_edge: int = 4096,
    ) -> str:
        ext = os.path.splitext(file_path)[1].lower()
        legacy_extractors = {
            '.pdf': self._extract_pdf_text,
            '.docx': self._extract_docx_text,
            '.txt': self._extract_plain_text,
            '.md': self._extract_plain_text,
            '.html': self._extract_html_text,
        }
        legacy = legacy_extractors.get(ext)
        if legacy is not None:
            return legacy(file_path)
        if not is_supported(ext, vision_enabled=vision_enabled):
            raise ValueError(f"Unsupported file type: {ext!r}")
        return extract_for_path(
            file_path,
            vision_enabled=vision_enabled,
            vision_max_bytes=vision_max_bytes,
            vision_max_edge=vision_max_edge,
        )

    def extract_metadata(self, file_path: str) -> Dict:
        ext = os.path.splitext(file_path)[1].lower()
        base = {
            'title': os.path.basename(file_path),
            'author': 'Unknown',
            'date': None,
            'file_type': ext,
        }
        if ext == '.pdf':
            base.update(self._pdf_metadata(file_path))
        elif ext == '.docx':
            base.update(self._docx_metadata(file_path))
        if base['date'] is None:
            base['date'] = datetime.now().isoformat()
        return base

    def clean_text(self, text: str) -> str:
        text = re.sub(r'\s+', ' ', text)
        return text.strip()

    def chunk_text(self, text: str) -> List[str]:
        token_ids = self._tokenizer.encode(text)
        if not token_ids:
            return []

        chunks: List[str] = []
        step = self.chunk_size - self.overlap
        start = 0
        while start < len(token_ids):
            end = min(start + self.chunk_size, len(token_ids))
            chunk = self._tokenizer.decode(token_ids[start:end])
            if chunk.strip():
                chunks.append(chunk)
            start += step
        return chunks

    def count_tokens(self, text: str) -> int:
        return len(self._tokenizer.encode(text))

    # --- private extractors ---

    def _extract_pdf_text(self, file_path: str) -> str:
        """Extract text page-by-page so one bad font map cannot abort the whole PDF."""
        with open(file_path, "rb") as handle:
            try:
                reader = PdfReader(handle)
            except PdfReadError as exc:
                raise ValueError(f"unreadable PDF: {file_path}") from exc
            parts: list[str] = []
            for index, page in enumerate(reader.pages):
                text = self._extract_one_pdf_page(file_path, index, page)
                parts.append(text)
            return "".join(parts)

    @staticmethod
    def _extract_one_pdf_page(file_path: str, index: int, page: Any) -> str:
        try:
            return page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001 — pypdf raises diverse font-map / cmap errors
            logger.warning(
                "pdf_page_extract_failed path=%s page=%s error=%s",
                file_path,
                index,
                exc,
            )
            return ""

    def _extract_docx_text(self, file_path: str) -> str:
        doc = Document(file_path)
        return '\n'.join(para.text for para in doc.paragraphs)

    def _extract_plain_text(self, file_path: str) -> str:
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()

    def _extract_html_text(self, file_path: str) -> str:
        with open(file_path, 'r', encoding='utf-8') as f:
            soup = BeautifulSoup(f, 'html.parser')
        return soup.get_text(separator=' ')

    # --- private metadata helpers ---

    def _pdf_metadata(self, file_path: str) -> Dict:
        result: Dict[str, Any] = {}
        try:
            with open(file_path, "rb") as handle:
                meta = PdfReader(handle).metadata
            if meta is None:
                return result
            title = getattr(meta, "title", None) or meta.get("/Title")
            author = getattr(meta, "author", None) or meta.get("/Author")
            created = getattr(meta, "creation_date", None) or meta.get("/CreationDate")
            if title:
                result["title"] = str(title)
            if author:
                result["author"] = str(author)
            if created:
                result["date"] = str(created)
        except (OSError, PdfReadError, TypeError, ValueError) as exc:
            logger.warning("pdf_metadata_failed path=%s error=%s", file_path, exc)
        return result

    def _docx_metadata(self, file_path: str) -> Dict:
        result = {}
        try:
            props = Document(file_path).core_properties
            if props.title:
                result['title'] = props.title
            if props.author:
                result['author'] = props.author
            if props.created:
                result['date'] = props.created.isoformat()
        except Exception:
            pass
        return result
