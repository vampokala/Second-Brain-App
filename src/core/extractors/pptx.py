"""PowerPoint .pptx → plain text (slide text + speaker notes)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation


def extract_pptx(path: Path | str) -> str:
    """Extract slide body and notes from a PowerPoint file."""
    path = Path(path)
    presentation = Presentation(str(path))
    sections: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines = [f"## Slide {index}"]
        body_parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                body_parts.append(shape.text.strip())
        if body_parts:
            lines.extend(body_parts)
        notes_slide = slide.notes_slide
        if notes_slide is not None:
            notes_text = notes_slide.notes_text_frame.text.strip()
            if notes_text:
                lines.append(f"Notes: {notes_text}")
        if len(lines) > 1:
            sections.append("\n".join(lines))
    return "\n\n".join(sections).strip()
