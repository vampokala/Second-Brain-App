"""CSV, TSV, and Excel → plain text for ingest."""

from __future__ import annotations

import csv
from pathlib import Path

_MAX_COLS = 32
_MAX_ROWS = 10_000


def _truncate_footer(truncated: bool) -> str:
    return "\n\n… [truncated]" if truncated else ""


def extract_csv(path: Path | str, *, max_rows: int = _MAX_ROWS) -> str:
    """Extract delimited text: tab-separated cells per row."""
    path = Path(path)
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    rows: list[str] = []
    truncated = False
    with path.open(newline="", encoding="utf-8", errors="replace") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for row_count, row in enumerate(reader):
            if row_count >= max_rows:
                truncated = True
                break
            cells = [cell.strip() for cell in row[:_MAX_COLS]]
            rows.append("\t".join(cells))
    return "".join(f"{line}\n" for line in rows) + _truncate_footer(truncated)


def extract_xlsx(path: Path | str, *, max_rows: int = _MAX_ROWS) -> str:
    """Extract Excel workbook sheets as markdown-style sections."""
    from openpyxl import load_workbook

    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[str] = []
    truncated = False
    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            section_lines = [f"## {sheet_name}", ""]
            rows_emitted = 0
            for row in sheet.iter_rows(values_only=True):
                if rows_emitted >= max_rows:
                    truncated = True
                    break
                cells = ["" if cell is None else str(cell).strip() for cell in row[:_MAX_COLS]]
                if not any(cells):
                    continue
                section_lines.append("\t".join(cells))
                rows_emitted += 1
            if len(section_lines) > 2:
                sections.append("\n".join(section_lines))
            if truncated:
                break
    finally:
        workbook.close()
    body = "\n\n".join(sections).strip()
    return body + _truncate_footer(truncated)
